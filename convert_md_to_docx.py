import os
import re
import sys
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def convert_md_to_docx(input_md, output_docx):
    doc = Document()
    
    # Define simple regex patterns for basic markdown
    # Bold: **text**
    bold_pattern = re.compile(r'\*\*(.*?)\*\*')
    # Italic: *text*
    italic_pattern = re.compile(r'\*(.*?)\*')
    # Inline code: `text`
    inline_code_pattern = re.compile(r'`(.*?)`')

    with open(input_md, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    
    for line in lines:
        line = line.strip('\n')
        
        # Handle code blocks
        if line.startswith('```'):
            in_code_block = not in_code_block
            continue
            
        if in_code_block:
            p = doc.add_paragraph(line)
            p.style = 'No Spacing'
            run = p.runs[0] if p.runs else p.add_run(line)
            run.font.name = 'Courier New'
            continue

        # Handle headings
        if line.startswith('# '):
            doc.add_heading(line[2:], level=0)
            continue
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=1)
            continue
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=2)
            continue
        elif line.startswith('#### '):
            doc.add_heading(line[5:], level=3)
            continue
        
        # Handle horizontal rule
        if line.strip() == '---':
            doc.add_page_break()
            continue
            
        # Handle lists
        is_list_item = False
        if line.startswith('- ') or line.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            line_content = line[2:]
            is_list_item = True
        elif line.strip() == '':
            continue
        else:
            p = doc.add_paragraph()
            line_content = line

        # Process inline styles (Bold, Italic, Code)
        # This is a simplified sequential processor
        
        # First, split by formatting markers to handle them correctly
        # For simplicity in this script, we'll just add the text and 
        # then apply some basic logic for bold/italic if they exist.
        # A more robust parser would be better, but for a one-off task:
        
        # Let's use a simple approach: add the run and then bold/italic components
        # (Handling nested formatting is complex, we'll do head-to-tail)
        
        remaining_text = line_content
        
        while remaining_text:
            # Check for bold, italic, or inline code
            b_match = bold_pattern.search(remaining_text)
            i_match = italic_pattern.search(remaining_text)
            c_match = inline_code_pattern.search(remaining_text)
            
            # Find the earliest match
            matches = []
            if b_match: matches.append((b_match.start(), 'bold', b_match))
            if i_match: matches.append((i_match.start(), 'italic', i_match))
            if c_match: matches.append((c_match.start(), 'code', c_match))
            
            if not matches:
                p.add_run(remaining_text)
                break
                
            matches.sort()
            start, m_type, match = matches[0]
            
            # Add text before the match
            if start > 0:
                p.add_run(remaining_text[:start])
                
            # Add the matched text with formatting
            content = match.group(1)
            run = p.add_run(content)
            if m_type == 'bold':
                run.bold = True
            elif m_type == 'italic':
                run.italic = True
            elif m_type == 'code':
                run.font.name = 'Courier New'
                
            # Update remaining text
            remaining_text = remaining_text[match.end():]

    doc.save(output_docx)
    print(f"Successfully converted {input_md} to {output_docx}")

if __name__ == "__main__":
    default_input = r"C:\Users\Khan1\OneDrive\Desktop\drive d data\PromoCode\user-stories-and-acceptance-criteria.md"
    default_output = r"C:\Users\Khan1\OneDrive\Desktop\drive d data\PromoCode\user-stories-and-acceptance-criteria.docx"

    input_file = sys.argv[1] if len(sys.argv) > 1 else default_input
    output_file = sys.argv[2] if len(sys.argv) > 2 else default_output

    if os.path.exists(input_file):
        convert_md_to_docx(input_file, output_file)
    else:
        print(f"Error: {input_file} not found.")
